from pptx import Presentation


def extract_pptx(file_path: str):

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        texts = []

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    texts.append(text)

        slide_text = "\n".join(texts)

        if slide_text:

            slides.append({
                "text": slide_text,
                "metadata": {
                    "source_type": "pptx",
                    "slide": slide_number
                }
            })

    return slides