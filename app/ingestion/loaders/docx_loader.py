from docx import Document
from io import BytesIO
import zipfile


def extract_docx_text(file_bytes):

    document = Document(
        BytesIO(file_bytes)
    )

    text_parts = []

    for paragraph in document.paragraphs:

        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract table content
    for table in document.tables:

        for row in table.rows:

            row_text = []

            for cell in row.cells:
                row_text.append(
                    cell.text.strip()
                )

            text_parts.append(
                " | ".join(row_text)
            )

    return "\n".join(text_parts)


def extract_docx_data(file_bytes):

    # 1. Extract text and tables
    text = extract_docx_text(file_bytes)

    # 2. Extract media (images)
    media = []

    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as z:

            for name in z.namelist():

                if name.startswith("word/media/"):

                    image_bytes = z.read(name)

                    ext = name.split(".")[-1].lower()

                    if ext in ["png", "jpg", "jpeg", "webp", "gif"]:

                        mime_type = f"image/{ext}" if ext != "jpg" else "image/jpeg"

                        media.append({
                            "type": "image",
                            "bytes": image_bytes,
                            "mime_type": mime_type,
                            "filename": name.split("/")[-1],
                            "metadata": {}
                        })

    except Exception as e:
        print(f"Error extracting images from DOCX: {e}")

    return text, media