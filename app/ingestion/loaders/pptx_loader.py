from pptx import Presentation
from io import BytesIO
import zipfile


def extract_pptx_text(file_bytes):

    presentation = Presentation(
        BytesIO(file_bytes)
    )

    text_parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):

        text_parts.append(
            f"Slide: {slide_number}"
        )

        for shape in slide.shapes:

            if hasattr(shape, "text"):

                text = shape.text.strip()

                if text:
                    text_parts.append(text)

    return "\n".join(text_parts)


def extract_pptx_data(file_bytes):

    # 1. Extract text content
    text = extract_pptx_text(file_bytes)

    # 2. Extract media (images)
    media = []

    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as z:

            for name in z.namelist():

                if name.startswith("ppt/media/"):

                    image_bytes = z.read(name)

                    ext = name.split(".")[-1].lower()

                    if ext in ["png", "jpg", "jpeg", "webp", "gif"]:

                        mime_type = f"image/{ext}" if ext != "jpg" else "image/jpeg"

                        media.append({
                            "type": "image",
                            "bytes": image_bytes,
                            "mime_type": mime_type,
                            "filename": name.split("/")[-1],
                            "metadata": {}
                        })

    except Exception as e:
        print(f"Error extracting images from PPTX: {e}")

    return text, media