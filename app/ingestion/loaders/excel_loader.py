import openpyxl
from io import BytesIO
import zipfile


def extract_excel_text(file_bytes):

    try:

        wb = openpyxl.load_workbook(
            BytesIO(file_bytes),
            data_only=True
        )

        text_parts = []

        for sheet_name in wb.sheetnames:

            sheet = wb[sheet_name]

            text_parts.append(
                f"Sheet: {sheet_name}"
            )

            sheet_data = []

            for row in sheet.iter_rows(
                values_only=True
            ):

                if any(
                    val is not None
                    for val in row
                ):

                    row_str = [
                        str(val).strip() if val is not None else ""
                        for val in row
                    ]

                    sheet_data.append(
                        " | ".join(row_str)
                    )

            if sheet_data:
                text_parts.append(
                    "\n".join(sheet_data)
                )

        wb.close()

        return "\n\n".join(text_parts)

    except Exception as e:

        print(
            f"Error reading excel with openpyxl: {e}"
        )

        return ""


def extract_excel_data(file_bytes):

    # 1. Extract text content
    text = extract_excel_text(file_bytes)

    # 2. Extract media (images)
    media = []

    try:
        with zipfile.ZipFile(BytesIO(file_bytes)) as z:

            for name in z.namelist():

                if name.startswith("xl/media/"):

                    image_bytes = z.read(name)

                    ext = name.split(".")[-1].lower()

                    if ext in ["png", "jpg", "jpeg", "webp", "gif"]:

                        mime_type = f"image/{ext}" if ext != "jpg" else "image/jpeg"

                        media.append({
                            "type": "image",
                            "bytes": image_bytes,
                            "mime_type": mime_type,
                            "filename": name.split("/")[-1],
                            "metadata": {}
                        })

    except Exception as e:
        print(f"Error extracting images from Excel: {e}")

    return text, media