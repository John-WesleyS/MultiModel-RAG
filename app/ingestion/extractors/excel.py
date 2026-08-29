from openpyxl import load_workbook


def extract_excel(file_path: str):

    workbook = load_workbook(
        file_path,
        data_only=True
    )

    documents = []

    for sheet in workbook.worksheets:

        rows = []

        for row in sheet.iter_rows(values_only=True):

            values = []

            for value in row:

                if value is not None:
                    values.append(str(value))

            if values:
                rows.append(" | ".join(values))

        text = "\n".join(rows)

        if text.strip():

            documents.append({
                "text": text,
                "metadata": {
                    "source_type": "excel",
                    "sheet": sheet.title
                }
            })

    workbook.close()

    return documents